import os
import csv


def _export_xlsx(rows, columns, outfile, sheet_title="Δεδομένα"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Δεδομένα"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2C5F8A", end_color="2C5F8A", fill_type="solid")
    for ci, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.font = header_font
        cell.fill = header_fill
    for ri, row in enumerate(rows, start=2):
        for ci, cell in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=cell)
    for ci, col in enumerate(columns, start=1):
        letter = get_column_letter(ci)
        ws.column_dimensions[letter].width = min(40, max(10, (col or "").__len__() + 8))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(outfile)
    return outfile


def _export_csv(rows, columns, outfile):
    with open(outfile, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(columns)
        for row in rows:
            writer.writerow(["" if v is None else v for v in row])
    return outfile


def _export_pdf(rows, columns, outfile, title="Δεδομένα", school_name="Βιβλιοθήκη Σχολείου"):
    from services.printing import table_pdf
    return table_pdf(rows, columns, outfile, title=title, school_name=school_name)


def _clean(rows, columns):
    out = []
    for row in rows:
        if row is None:
            continue
        if isinstance(row, dict):
            vals = []
            for c in columns:
                v = row.get(c, "")
                if isinstance(v, (list, tuple)):
                    v = ", ".join(str(x) for x in v)
                vals.append(v)
            out.append(vals)
        else:
            out.append(row)
    return out


def export_rows(rows, columns, outfile, fmt="xlsx", title="Δεδομένα", school_name="Βιβλιοθήκη Σχολείου"):
    rows = _clean(rows, columns)
    fmt = (fmt or "xlsx").lower().lstrip(".")
    if fmt in ("xlsx", "xls"):
        return _export_xlsx(rows, columns, outfile, sheet_title=title)
    if fmt == "csv":
        return _export_csv(rows, columns, outfile)
    if fmt == "pdf":
        return _export_pdf(rows, columns, outfile, title=title, school_name=school_name)
    raise ValueError(f"Μη υποστηριζόμενη μορφή: {fmt}")


def autodetect_format(path):
    ext = os.path.splitext(path or "")[1].lower()
    if ext == ".xlsx":
        return "xlsx"
    if ext == ".xls":
        return "xls"
    if ext == ".pdf":
        return "pdf"
    return "csv"