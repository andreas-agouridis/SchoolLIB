import os
import datetime
import re

BOOK_FIELDS = ["book_number", "title", "author", "publisher", "edition", "year",
               "pages", "isbn", "dewey", "copies", "notes", "categories", "tags", "shelves"]
MEMBER_FIELDS = ["member_number", "name", "class", "email", "phone", "notes"]
LOAN_FIELDS = ["member_number", "book_number", "loan_date", "due_date", "return_date", "renewals"]

SYNONYMS = {
    "title": ["title", "titulo", "ober title", "titolo", "κaaτικ", "τίτλος", "τiτλος", "τιτλος"],
    "author": ["author", "creators", "pelaut", "authors", "autor", "συγγραφέας", "συγγραφεας", "συγγραφεασ"],
    "isbn": ["isbn", "ean", "barcode"],
    "book_number": ["book number", "book_number", "book no", "accno", "accession", "accession number",
                    "inclination", "barcode", "αρ βιβλίου", "αριθμός βιβλίου", "αριθμοσ βιβλίου",
                    "αρ. βιβλιου", "αρ.", "αριθμός καταχώρησης", "κωδικός", "codigo", "id"],
    "publisher": ["publisher", "editorial", "editor", "εκδότης", "εκδοτης", "εκδοτικοσ οικοσ", "imprenta"],
    "edition": ["edition", "έκδοση", "εκδοση", "edicao"],
    "year": ["year", "publish date", "publication year", "data publi", "χρονολογία", "ετος", "έτος", "año", "ano"],
    "pages": ["pages", "paginas", "page count", "σελίδες", "σελιδεσ", "no of pages"],
    "dewey": ["dewey", "ταξιθετικός", "ταξιθετικος", "ταξιθετικός αριθμός", "call number"],
    "copies": ["copies", "quantity", "αντίτυπα", "αντιτυπα", "num copies"],
    "notes": ["notes", "σημειώσεις", "σημειωσεις", "remarks"],
    "categories": ["category", "categories", "genres", "subject", "subjects", "thema", "θέμα", "κατηγορία",
                   "κατηγορια", "κaτηγoριεσ", "topics"],
    "tags": ["tags", "keywords", "λέξεις κλειδιά", "κλειδια"],
    "shelves": ["location", "shelf", "shelves", "location 2", "topográfica", "τοποθεσία", "τοποθεσια", "ράφι",
                "ραφι", "τοποθεσια ραφιου"],
    "member_number": ["member number", "member_number", "member no", "card no", "memberid", "reader no",
                      "αρ μέλους", "αριθμός μέλους", "αριθμοσ μέλους", "μέλος", "κωδικός μέλους", "no. socio"],
    "name": ["name", "full name", "reader", "reader name", "member", "ονοματεπώνυμο", "ονομα", "όνομα",
             "ονοματεπωνυμο", "ονoμα", "cliente name"],
    "class": ["class", "grade", "form", "classe", "curso", "τάξη", "ταξη", "τaξη"],
    "email": ["email", "e-mail", "correo", "ηλ. ταχυδρομείο", "μαιλ"],
    "phone": ["phone", "telephone", "mobile", "τηλέφωνο", "τηλεφωνο", "telefon"],
    "loan_date": ["loan date", "date issue", "issue date", "checkout date", "ημερομηνία δανεισμού",
                  "ημερομηνια δανεισμου", "δανείστηκε", "fecha prestamo"],
    "due_date": ["due date", "date due", "return due", "προθεσμία", "προθεσμια", "ημερομηνία επιστροφής",
                 "ημερομηνια επιστροφης", "fecha vencimiento"],
    "return_date": ["return date", "date returned", "date of return", "επιστράφηκε", "ημερομηνία επιστροφής",
                    "ημερομηνια επιστροφης"],
    "renewals": ["renewals", "renewed", "ανανεώσεις", "ανανεωσεις"],
}


def normalize_header(h):
    h = re.sub(r"\s+", " ", (h or "").strip().lower())
    h = re.sub(r"[*|/\\:;.,_-]+", " ", h)
    return h.strip()


def build_synonym_map():
    m = {}
    for field, words in SYNONYMS.items():
        for w in words:
            key = normalize_header(w)
            if key:
                m[key] = field
    return m


_SYN = build_synonym_map()


def guess_field(header):
    h = normalize_header(header)
    if h in _SYN:
        return _SYN[h]
    for key, field in _SYN.items():
        if key in h or h in key:
            return field
    if "isbn" in h or "ean" in h:
        return "isbn"
    if "titl" in h:
        return "title"
    if "aut" in h or "creat" in h:
        return "author"
    if "publish" in h:
        return "publisher"
    if "nomb" in h or "name" in h:
        return "name"
    return None


def _read_csv(path):
    import csv
    import io
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1253", "cp1250", "cp1252", "latin-1"):
        try:
            with open(path, encoding=enc, newline="") as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Δεν ήταν δυνατή η ανάγνωση του CSV (κωδικοποίηση). Εξάγετε σε .xlsx.")
    lines = text.splitlines()
    first = lines[0] if lines else ""
    delims = {",": first.count(","), ";": first.count(";"), "\t": first.count("\t")}
    delim = max(delims, key=delims.get)
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    return rows


def read_rows_file(path):
    ext = os.path.splitext(path)[1].lower()
    headers = []
    rows = []
    if ext in (".xlsx", ".xls"):
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            wb.close()
        else:
            import xlrd
            book = xlrd.open_workbook(path)
            ws = book.sheet_by_index(0)
            rows = [ws.row_values(r) for r in range(ws.nrows)]
    else:
        rows = _read_csv(path)
    if rows:
        headers = [str(h).strip() for h in rows[0]]
        rows = rows[1:]
    data = []
    for r in rows:
        if r is None:
            continue
        d = {}
        for i, h in enumerate(headers):
            val = r[i] if i < len(r) else None
            if isinstance(val, float) and val.is_integer():
                val = int(val)
            d[h] = "" if val is None else str(val).strip()
        if any(v for v in d.values()):
            data.append(d)
    return {"headers": headers, "rows": data}


def guess_mapping(headers, fields, required=None):
    mapping = {}
    required = required or []
    candidates = {f: [] for f in fields}
    for h in headers:
        g = guess_field(h)
        if g in candidates:
            candidates[g].append(h)
    for f in fields:
        if candidates[f]:
            mapping[f] = candidates[f][0]
    return mapping, candidates


def _as_int(v, fallback=0):
    if v is None:
        return fallback
    s = str(v).strip()
    try:
        m = re.search(r"\d+", s)
        if not m:
            return fallback
        return int(m.group())
    except Exception:
        return fallback


def _as_date(v):
    if not v:
        return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%m/%d/%Y"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except Exception:
            continue
    return None


def _split_list(v):
    if not v:
        return []
    return [x.strip() for x in re.split(r"[,;|/]", str(v)) if x.strip()]


def import_books(db, rows, mapping, skip_duplicates=True, update_existing=False, library_id=None):
    report = {"imported": 0, "updated": 0, "skipped": 0, "errors": []}
    cat_cache = {}
    for row in rows:
        title = (row.get(mapping.get("title")) or "").strip()
        if not title:
            report["skipped"] += 1
            report["errors"].append(f"Χωρίς τίτλο: {len(report['errors'])+1}")
            continue
        isbn = (row.get(mapping.get("isbn")) or "").strip()
        dup = None
        if skip_duplicates or update_existing:
            dup = db.find_duplicate(title, isbn)
        if dup and not update_existing:
            report["skipped"] += 1
            report["errors"].append(f"Διπλότυπο: {title}")
            continue
        cats = []
        for chunk in _split_list(row.get(mapping.get("categories")) if mapping.get("categories") else ""):
            if chunk not in cat_cache:
                res = db.add_category(chunk)
                cat_cache[chunk] = res.get("id")
            if cat_cache.get(chunk):
                cats.append(cat_cache[chunk])
        tags = _split_list(row.get(mapping.get("tags")) if mapping.get("tags") else "")
        shelves = _split_list(row.get(mapping.get("shelves")) if mapping.get("shelves") else "")
        shelf_ids = []
        for sname in shelves:
            sid = _find_or_create_shelf(db, sname, library_id)
            if sid:
                shelf_ids.append(sid)
        data = {
            "library_id": library_id,
            "book_number": row.get(mapping["book_number"]) if mapping.get("book_number") else "",
            "title": title,
            "author": row.get(mapping["author"]) if mapping.get("author") else "",
            "publisher": row.get(mapping["publisher"]) if mapping.get("publisher") else "",
            "edition": row.get(mapping["edition"]) if mapping.get("edition") else "",
            "year": row.get(mapping["year"]) if mapping.get("year") else "",
            "pages": _as_int(row.get(mapping["pages"]) if mapping.get("pages") else ""),
            "isbn": isbn,
            "dewey": row.get(mapping["dewey"]) if mapping.get("dewey") else "",
            "copies": _as_int(row.get(mapping["copies"]) if mapping.get("copies") else "", 1) or 1,
            "categories": cats,
            "tags": tags,
            "shelves": shelf_ids,
        }
        if dup and update_existing:
            db.update_book(dup["id"], data)
            report["updated"] += 1
        else:
            res = db.add_book(data)
            if res.get("ok"):
                report["imported"] += 1
            else:
                report["skipped"] += 1
                report["errors"].append(f"{title}: {res.get('error')}")
    return report


def _find_or_create_shelf(db, name, library_id=None):
    name = (name or "").strip()
    if not name:
        return None
    parts = [p.strip() for p in name.split(">") if p.strip()]
    parent = None
    libs = db.list_libraries()
    lib_id = library_id or (libs[0]["id"] if libs else 1)
    existing = db.list_shelves(lib_id)
    for depth, part in enumerate(parts):
        found = None
        for s in existing:
            if s["name"] == part and (s["parent_id"] == parent or (parent is None and s["parent_id"] is None)):
                found = s
                break
        if found:
            parent = found["id"]
        else:
            sid = db.add_shelf(part, parent_id=parent, library_id=lib_id)
            existing = db.list_shelves(lib_id)
            parent = sid
    return parent


def import_members(db, rows, mapping, skip_duplicates=True, library_id=None):
    report = {"imported": 0, "updated": 0, "skipped": 0, "errors": []}
    for row in rows:
        name = (row.get(mapping.get("name")) or "").strip()
        if not name:
            report["skipped"] += 1
            report["errors"].append("Χωρίς όνομα μέλους")
            continue
        num = (row.get(mapping["member_number"]) if mapping.get("member_number") else "").strip()
        existing = None
        for m in db.list_members(q=name):
            if m["name"].strip().lower() == name.lower():
                existing = m
                break
        data = {
            "library_id": library_id,
            "member_number": num,
            "name": name,
            "class": row.get(mapping["class"]) if mapping.get("class") else "",
            "email": row.get(mapping["email"]) if mapping.get("email") else "",
            "phone": row.get(mapping["phone"]) if mapping.get("phone") else "",
        }
        if existing and not skip_duplicates:
            db.update_member(existing["id"], data)
            report["updated"] += 1
        elif existing:
            report["skipped"] += 1
            report["errors"].append(f"Ήδη υπάρχει: {name}")
        else:
            db.add_member(data)
            report["imported"] += 1
    return report


def import_loans(db, rows, mapping, due_days=14):
    report = {"imported": 0, "skipped": 0, "errors": []}
    for row in rows:
        mnum = (row.get(mapping.get("member_number")) or "").strip()
        bnum = (row.get(mapping.get("book_number")) or "").strip()
        if not mnum and not bnum:
            report["skipped"] += 1
            report["errors"].append("Κενό βιβλίο/μέλος")
            continue
        mem = None
        for m in db.list_members(q=mnum):
            if m["member_number"] == mnum:
                mem = m
                break
        if not mem:
            report["skipped"] += 1
            report["errors"].append(f"Μέλος δεν βρέθηκε: {mnum}")
            continue
        book = None
        for b in db.search_books(bnum):
            if b["book_number"] == bnum or b["isbn"] == bnum:
                book = b
                break
        if not book:
            report["skipped"] += 1
            report["errors"].append(f"Βιβλίο δεν βρέθηκε: {bnum}")
            continue
        loan_date = _as_date(row.get(mapping.get("loan_date")) if mapping.get("loan_date") else "")
        due_date = _as_date(row.get(mapping.get("due_date")) if mapping.get("due_date") else "")
        return_date = _as_date(row.get(mapping.get("return_date")) if mapping.get("return_date") else "")
        if not loan_date:
            loan_date = datetime.date.today()
        if not due_date:
            due_date = loan_date + datetime.timedelta(days=due_days)
        if return_date and return_date < loan_date:
            report["skipped"] += 1
            report["errors"].append(f"{bnum}: ημερομηνία επιστροφής πριν τον δανεισμό")
            continue
        res = db.add_loan_record(book["id"], mem["id"], loan_date, due_date,
                                 return_date=return_date)
        if not res.get("ok"):
            report["skipped"] += 1
            report["errors"].append(f"{bnum}: {res.get('error')}")
            continue
        report["imported"] += 1
    return report


def export_template(path, kind="books"):
    from services import exporter
    if kind == "books":
        headers = ["book_number", "title", "author", "publisher", "edition", "year",
                   "pages", "isbn", "dewey", "copies", "categories", "tags", "shelves"]
        rows = [["B001", "Το Δέντρο", "Νίκος Παπαδόπουλος", "Εκδόσεις Αθήνα", "1η", "2020",
                 120, "9789600000001", "889", 2, "Λογοτεχνία", "ελληνική, μυθιστόρημα", "Ράφι 1"]]
    elif kind == "members":
        headers = ["member_number", "name", "class", "email", "phone", "notes"]
        rows = [["M001", "Γιώργος Παπαδόπουλος", "Γ1", "g@example.gr", "2100000000", ""]]
    else:
        headers = ["member_number", "book_number", "loan_date", "due_date", "return_date"]
        rows = [["M001", "B001", "2025-01-10", "2025-01-24", ""]]
    return exporter.export_rows(rows, headers, os.fspath(path) if path else "template.xlsx", fmt="xlsx")